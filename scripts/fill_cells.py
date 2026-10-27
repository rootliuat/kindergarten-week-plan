# -*- coding: utf-8 -*-
"""
[INPUT]: 依赖 python-docx 读取输入 DOCX，依赖 spec JSON 指定目标内容单元格
[OUTPUT]: 输出另存的已填写 DOCX，拒绝覆盖原件、修改首页或访问越界单元格
[POS]: scripts 模块的安全写入器，与 inspect_doc.py 的表格索引结果配合
[PROTOCOL]: 变更时更新此头部，然后检查 AGENTS.md

周计划填表器：只往内容格填字，绝不改行高列宽。
用法: python fill_cells.py <输入docx> <spec.json> <输出docx>
spec.json: [{"table":1,"row":1,"lines":["第一段","第二段"]}, ...]
格式硬指标：仿宋_GB2312 小四(12pt)、行距固定值22磅。"""
import sys, io, json
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

FONT = "仿宋_GB2312"

def style_run(r):
    r.font.name = FONT
    r.font.size = Pt(12)
    rPr = r._element.get_or_add_rPr()
    rf = rPr.get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), FONT)
    rf.set(qn('w:ascii'), FONT)
    rf.set(qn('w:hAnsi'), FONT)

def fill(cell, lines):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    for i, line in enumerate(lines):
        p = p0 if i == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = Pt(22)   # 固定值22磅
        r = p.add_run(line)
        style_run(r)

def target_cell(document, item):
    table_index = item["table"]
    row_index = item["row"]
    col_index = item.get("col", 1)
    if table_index == 0:
        raise ValueError("禁止修改首页 table 0")
    try:
        table = document.tables[table_index]
        row = table.rows[row_index]
        return row.cells[col_index]
    except IndexError as exc:
        raise ValueError(
            f"单元格越界: table={table_index}, row={row_index}, col={col_index}"
        ) from exc

def main(docx_path, spec_path, output_path):
    source = Path(docx_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise ValueError("输出路径必须不同于输入路径，禁止覆盖原件")
    with open(spec_path, encoding='utf-8') as f:
        spec = json.load(f)
    d = Document(docx_path)
    for item in spec:
        cell = target_cell(d, item)
        fill(cell, item["lines"])
        print(f"filled table {item['table']} row {item['row']}: {len(item['lines'])}段")
    output.parent.mkdir(parents=True, exist_ok=True)
    d.save(output_path)
    print("saved:", output_path)

if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("用法: python fill_cells.py <输入docx> <spec.json> <输出docx>")
    main(sys.argv[1], sys.argv[2], sys.argv[3])
