# -*- coding: utf-8 -*-
"""
[INPUT]: 依赖 python-docx 读取幼儿园周计划 DOCX 的正文与表格
[OUTPUT]: 输出表头、日期、表格布局、上午下午类型及模板完整性状态
[POS]: scripts 模块的只读侦察器，为内容映射和安全写入提供结构索引
[PROTOCOL]: 变更时更新此头部，然后检查 AGENTS.md

周计划 DOCX 结构检查器。用法: python inspect_doc.py <docx路径>
"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

def main(path):
    d = Document(path)
    all_text = "\n".join(p.text for p in d.paragraphs)
    headers = [p.text.strip() for p in d.paragraphs if "周计划表" in p.text]
    for header in headers:
        print("HEADER:", header)
    dates = re.findall(r"\d+\s*月\s*\d+\s*日\s*星期[一二三四五六日天]\s*(?:上午|下午)?", all_text)
    for date in dates:
        print("DATE:", re.sub(r"\s+", " ", date).strip())
    print("tables:", len(d.tables))
    morning = afternoon = 0
    for ti, tb in enumerate(d.tables):
        print(f"== table {ti}: {len(tb.rows)}x{len(tb.columns)}")
        labels = "|".join(cell.text.replace("\n", "") for row in tb.rows for cell in row.cells)
        if "活动目标" in labels and "活动过程" in labels:
            morning += 1
            print("  TYPE: morning")
        elif "指导要点" in labels and "活动回顾" in labels:
            afternoon += 1
            print("  TYPE: afternoon")
        for ri, row in enumerate(tb.rows):
            cells = [c.text.replace("\n", "/").strip() for c in row.cells]
            # 合并单元格会重复文本，压掉连续重复
            compact = [cells[0]] + [t for i, t in enumerate(cells[1:], 1) if t != cells[i-1]]
            line = " | ".join(compact)
            print(f"  R{ri}: {line[:120]}")
    if len(d.tables) <= 1:
        status = "homepage-only"
    elif morning and afternoon:
        status = "has-both-templates"
    elif morning:
        status = "morning-only"
    elif afternoon:
        status = "afternoon-only"
    else:
        status = "unknown-layout"
    print(f"TEMPLATE_STATUS: {status}; morning={morning}; afternoon={afternoon}")

if __name__ == "__main__":
    main(sys.argv[1])
